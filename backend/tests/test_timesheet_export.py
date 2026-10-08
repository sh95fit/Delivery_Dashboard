import io
from datetime import date, time

from openpyxl import load_workbook

from app.services.timesheet_excel import schedule_ranges
from app.services.timesheet_export import DAY_COL, FIRST_ROW, build_timesheet, schedule_note, verify_timesheet

M = date(2026, 10, 1)   # 31일 → E..AI 날짜, AJ 월, AL 월급여, AM 인센, AN 주말, AO 총지급액, AP 부가세


def P(name, rate=16000, days=None, **kw):
    return {"name": name, "rate": rate, "note": "07:00~12:30", "days": days or {}, "pay_value": None,
            "incentive": 0, "weekend": 0, "vat": False, **kw}


A = P("인력A", days={1: {"in": 420, "out": 793, "break": 0}, 2: {"in": 430, "out": 750, "break": 30}})
B = P("인력B", 27500, days={5: {"in": 390, "out": 838, "break": 0}}, vat=True, incentive=10000)


def ws_of(data):
    return load_workbook(io.BytesIO(data)).active


def test_build_and_verify_roundtrip():
    for sheet in ("labor", "business"):
        people = [A, B, P("인력C")]
        data = build_timesheet(sheet, M, people, factory="테스트공장", dept="물류팀")
        assert verify_timesheet(data, people) == []


def test_formulas_and_values():
    ws = ws_of(build_timesheet("business", M, [A, B]))
    assert ws.cell(FIRST_ROW + 3, DAY_COL).value == '=IF(COUNT(E8:E9)=2,E9-E8-N(E10),"")'
    assert ws["F10"].value == time(0, 30)
    assert ws["AJ11"].value == "=SUM(E11:AI11)"
    assert ws["AL8"].value == "=AJ11*24"
    assert ws["AL11"].value == "=ROUND($B8*AL8,0)"
    assert ws["AO11"].value == "=SUM(AL11:AN11)"
    assert ws["AP11"].value is None                     # 부가세 없음
    assert ws["AM15"].value == 10000 and ws["AP15"].value == "=ROUND(AO15*1.1,0)"
    assert ws["A16"].value == "총계"
    assert ws["AO16"].value == '=SUMIF($D$8:$D$15,"합계",AO8:AO15)'
    assert ws_of(build_timesheet("labor", M, [A]))["AP5"].value is None   # 근로소득은 부가세 열 없음


def test_tampered_file_detected():
    wb = load_workbook(io.BytesIO(build_timesheet("labor", M, [A])))
    wb.active["E9"].value = time(13, 30)               # 1일 퇴근 변경
    buf = io.BytesIO()
    wb.save(buf)
    assert any("1일" in e for e in verify_timesheet(buf.getvalue(), [A]))


def test_people_mismatch_detected():
    assert verify_timesheet(build_timesheet("labor", M, [A]), [A, P("인력C")])


def test_pay_value_when_rate_changed():
    assert ws_of(build_timesheet("labor", M, [{**A, "pay_value": 123456}]))["AL11"].value == 123456


def test_thirty_day_month():
    data = build_timesheet("labor", date(2026, 9, 1), [A])
    assert verify_timesheet(data, [A]) == []
    assert ws_of(data)["AI5"].value == "월"


def test_schedule_note():
    s = schedule_note([(date(2026, 9, 1), "07:00", "12:30", 16000), (date(2026, 9, 8), "07:00", "13:30", 16000)])
    assert s == "07:00~12:30 / 9/8~ 07:00~13:30"
    assert schedule_ranges(s) == [("07:00", "12:30"), ("07:00", "13:30")]
    assert schedule_note([(date(2026, 9, 1), "07:00", "12:30", 16000),
                          (date(2026, 9, 15), "07:00", "12:30", 17000)]) == "07:00~12:30 / 9/15~ 시급 17,000"
