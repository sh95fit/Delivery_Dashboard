from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.services.excel_export import build_xlsx
from app.services.fuel_excel import candidates, judge, parse_file

HEAD = ["거래처", "차량", "순번", "제품명", "일자", "수량", "공급가액", "부가세", "합계"]
CO = "111111 / (주)런치랩"
TODAY = date(2026, 10, 16)
VEH = [
    {"id": 1, "plate_no": "802더1612", "fuel_type": "diesel", "vat_deductible": False},
    {"id": 2, "plate_no": "80가3300", "fuel_type": "diesel", "vat_deductible": True},
    {"id": 3, "plate_no": "12나 3300", "fuel_type": "diesel", "vat_deductible": True},
    {"id": 4, "plate_no": "33하9426", "fuel_type": "diesel", "vat_deductible": False},
]


def html(rows, enc="cp949") -> bytes:
    body = "".join("<tr>" + "".join(f"<td>{'' if c is None else c}</td>" for c in r) + "</tr>" for r in rows)
    return f"<html><body><table>{body}</table></body></html>".encode(enc)


def R(code, d, total, supply=None, vat=None, product="저유황경유", qty="50"):
    return [CO, code, "00", product, d, qty, supply, vat, total]


SAMPLE = [
    HEAD,
    R("1612", "20260804", "152000", "138182", "13818", qty="80.551"),
    R("1612", "20260814", "50000", "45455", "4545", qty="26.497"),
    [None, "1612 합계", None, None, None, "107.048", "183637", "18363", "202000"],
    R("9426", "20260804", "61000", "55455", "5545", product="무연휘발유", qty="32.156"),
    [None, "9426 합계", None, None, None, "32.156", "55455", "5545", "61000"],
    ["(주)런치랩 합계", None, None, None, None, "139.204", "239092", "23908", "263000"],
    ["-------------------------------", None, None, None, None, "0", "0", "0", "0"],
    ["[총 합계]", None, None, None, None, "139.204", "239092", "23908", "263000"],
    [None, None, None, None, None, "0", "0", "0", "0"],
]


def items(rows):
    return parse_file(html([HEAD, *rows]))["items"]


def test_parse_html_sample():
    p = parse_file(html(SAMPLE))
    assert p["file_kind"] == "html" and p["header_row"] == 1
    assert [(i["row"], i["code"], i["date"], i["total"], i["fuel_type"]) for i in p["items"]] == [
        (2, "1612", date(2026, 8, 4), 152000, "diesel"),
        (3, "1612", date(2026, 8, 14), 50000, "diesel"),
        (5, "9426", date(2026, 8, 4), 61000, "gasoline"),
    ]
    assert p["skipped"] == 6
    assert not [c for c in p["checks"] if c["level"] == "warn"]
    assert any("263,000" in c["text"] for c in p["checks"])


def test_parse_utf8_and_subtotal_mismatch():
    rows = [r[:] for r in SAMPLE]
    rows[3][8] = "999"
    p = parse_file(html(rows, "utf-8"))
    assert any(c["level"] == "warn" and "1612" in c["text"] for c in p["checks"])


def test_xlsx_numbers():
    wb = Workbook()
    ws = wb.active
    ws.append(HEAD)
    ws.append([CO, 1612, "00", "저유황경유", 20260804, 80.551, 138182, 13818, 152000])
    ws.append([None, "1612 합계", None, None, None, 80.551, 138182, 13818, 152000])
    buf = BytesIO()
    wb.save(buf)
    p = parse_file(buf.getvalue())
    assert p["file_kind"] == "xlsx" and len(p["items"]) == 1
    it = p["items"][0]
    assert (it["code"], it["date"], it["total"], it["supply"]) == ("1612", date(2026, 8, 4), 152000, 138182)


def test_not_excel_and_no_header():
    with pytest.raises(ValueError):
        parse_file(b"hello")
    with pytest.raises(ValueError):
        parse_file(html([["a", "b"], ["1", "2"]]))


def test_candidates():
    assert [v["id"] for v in candidates("1612", VEH)] == [1]
    assert sorted(v["id"] for v in candidates("3300", VEH)) == [2, 3]
    assert candidates("0000", VEH) == []
    assert [v["id"] for v in candidates("802더 1612", VEH)] == [1]


def test_judge_auto_and_manual():
    rows = [R("1612", "20260901", "144000", "130909", "13091"), R("3300", "20260902", "70000"),
            R("0000", "20260903", "50000")]
    a, b, c = judge(items(rows), VEH, {}, set(), set(), TODAY)
    assert a["level"] == "ok" and a["vehicle_id"] == 1 and a["auto"] and a["checked"] and a["cost"] == 144000
    assert b["level"] == "need_vehicle" and sorted(b["candidates"]) == [2, 3] and not b["selectable"]
    assert c["level"] == "need_vehicle" and c["candidates"] == []
    _, b2, c2 = judge(items(rows), VEH, {"3300": 2, "0000": 1}, set(), set(), TODAY)
    assert b2["vehicle_id"] == 2 and b2["checked"] and (b2["supply_amount"], b2["vat_amount"]) == (63636, 6364)
    assert b2["cost"] == 63636            # 공제 차량 → 공급가액
    assert c2["vehicle_id"] == 1 and c2["selectable"]


def test_judge_warnings():
    rows = [R("9426", "20260904", "59000", product="무연휘발유"), R("1612", "20260901", "144000"),
            R("1612", "20260903", "80000"), R("1612", "20260903", "80000"),
            R("1612", "20260905", "100000", "90000", "9000")]
    fuel, dup_db, first, dup_file, amt = judge(items(rows), VEH, {}, {(1, date(2026, 9, 1), 144000)}, set(), TODAY)
    assert fuel["level"] == "warn" and fuel["selectable"] and not fuel["checked"]
    assert "이미 있습니다" in dup_db["messages"][0] and not dup_db["checked"]
    assert first["checked"] and not dup_file["checked"] and "파일 안에" in dup_file["messages"][0]
    assert amt["checked"] and (amt["supply_amount"], amt["vat_amount"]) == (90909, 9091)


def test_judge_errors():
    rows = [R("1612", "20261020", "50000"), R("1612", "20260815", "50000"),
            R("1612", "2026-13-01", "50000"), R("1612", "20260910", "-50000")]
    out = judge(items(rows), VEH, {}, set(), {date(2026, 8, 1)}, TODAY)
    assert [r["level"] for r in out] == ["error"] * 4
    assert "미래" in out[0]["messages"][0] and "마감" in out[1]["messages"][0]
    assert "날짜" in out[2]["messages"][0] and "취소" in out[3]["messages"][0]
    assert not any(r["selectable"] for r in out)


def test_refund_pair():
    rows = [R("1612", "20260910", "50000"), R("1612", "20260910", "-50000"), R("1612", "20260912", "60000")]
    out = judge(items(rows), VEH, {}, set(), set(), TODAY)
    assert [r["level"] for r in out] == ["cancel", "cancel", "ok"]
    assert (out[0]["pair"], out[1]["pair"]) == (3, 2)


def test_build_xlsx():
    body = build_xlsx([{"title": "주유 기록", "columns": [("주유일자", "date", "date"), ("합계", "total", "won")],
                        "rows": [{"date": date(2026, 9, 1), "total": 144000}],
                        "totals": {"date": "합계", "total": 144000}}])
    ws = load_workbook(BytesIO(body)).active
    assert ws.title == "주유 기록" and ws["A1"].value == "주유일자"
    assert ws["B2"].value == 144000 and ws["B2"].number_format == "#,##0"
    assert ws["A3"].value == "합계" and ws.freeze_panes == "A2"
