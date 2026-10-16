"""S3-P1b-2b-3 주유 거래내역 엑셀 읽기 · 행 판정 (DB 없음 → 단위 테스트 대상).

양식 (주유소 월 거래내역, 2026-08·09 파일 기준)
  거래처 | 차량(번호 끝 4자리) | 순번 | 제품명 | 일자(20260804) | 수량 | 공급가액 | 부가세 | 합계
- 'NNNN 합계' · '(주)런치랩 합계' · '----' · '[총 합계]' · 0만 있는 줄은 저장하지 않고 합계 줄로 검산만 한다
- 열은 위치가 아니라 머리글 이름으로 찾는다 → 열 순서가 바뀌거나 열이 늘어도 동작
- .xlsx · 구형 .xls · 확장자만 xls인 HTML 표를 파일 앞부분으로 판별해 모두 읽는다
- 차량: 번호 전체가 같으면 그 차량, 아니면 끝 4자리가 한 대에만 맞을 때 자동. 0대·2대 이상이면 직접 선택
- 취소(음수) 행은 같은 차량·같은 금액 원 거래와 짝지어 둘 다 저장하지 않는다
"""
from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from html.parser import HTMLParser
from io import BytesIO

from app.services.fuel_log_calc import amounts

MAX_ROWS = 2000
HEADER_SCAN = 20
EXCEL_EPOCH = date(1899, 12, 30)
XLS_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")
FUEL_LABEL = {"gasoline": "휘발유", "diesel": "경유", "lpg": "LPG", "ev": "전기"}

# 머리글 후보 (공백 무시, 이름이 완전히 같을 때만). 앞에 있는 이름이 우선
COLS = {
    "date": ("일자", "거래일자", "거래일", "주유일자", "주유일", "이용일자", "이용일", "승인일자", "승인일"),
    "plate": ("차량", "차량번호", "차번"),
    "total": ("합계", "합계금액", "이용금액", "결제금액", "승인금액", "금액"),
    "supply": ("공급가액", "공급가"),
    "vat": ("부가세", "세액"),
    "qty": ("수량", "주유량", "리터"),
    "product": ("제품명", "유종", "품명", "유류"),
}
REQUIRED = ("date", "plate", "total")


class NoHeader(ValueError):
    pass


# ---------- 값 정리 ----------
def _cell(v):
    if isinstance(v, str):
        s = v.replace("\xa0", " ").strip()
        return s or None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _txt(v) -> str:
    v = _cell(v)
    return "" if v is None else str(v).strip()


def _key(v) -> str:
    return re.sub(r"\s+", "", _txt(v))


def _num(v) -> float | None:
    v = _cell(v)
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("원", "").strip()
    if s.startswith("(") and s.endswith(")"):
        s = "-" + s[1:-1]
    try:
        return float(s)
    except ValueError:
        return None


def _won(v) -> int | None:
    n = _num(v)
    return None if n is None else int(round(n))


def parse_date(v) -> date | None:
    v = _cell(v)
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        n = int(v)
        if 20000 <= n <= 80000:          # 엑셀 날짜 일련번호
            return EXCEL_EPOCH + timedelta(days=n)
        v = str(n)
    s = str(v).strip()
    m = re.match(r"^(\d{4})(\d{2})(\d{2})$", s) or re.match(r"^(\d{4})\s*[-./]\s*(\d{1,2})\s*[-./]\s*(\d{1,2})", s)
    if not m:
        return None
    try:
        return date(int(m[1]), int(m[2]), int(m[3]))
    except ValueError:
        return None


def fuel_kind(product) -> str | None:
    s = _key(product).lower()
    if not s:
        return None
    if "경유" in s or "디젤" in s or "diesel" in s:
        return "diesel"
    if "휘발유" in s or "가솔린" in s or "gasoline" in s:
        return "gasoline"
    if "lpg" in s or "부탄" in s or "프로판" in s:
        return "lpg"
    if "전기" in s or "충전" in s or "kwh" in s:
        return "ev"
    return None


def norm_code(v) -> str:
    """차량 칸 → 비교용 코드. 공백·하이픈 제거, 숫자만 4자리 미만이면 앞을 0으로 (엑셀이 0123을 123으로 바꾼 경우)."""
    s = re.sub(r"[\s-]", "", _txt(v))
    return s.zfill(4) if s.isdigit() and len(s) < 4 else s


def _plate(p) -> str:
    return re.sub(r"[\s-]", "", str(p or ""))


def candidates(code: str, vehicles: list[dict]) -> list[dict]:
    """번호 전체가 같은 차량 → 없으면 끝 4자리 이상이 맞는 차량들."""
    c = norm_code(code)
    if not c:
        return []
    exact = [v for v in vehicles if _plate(v["plate_no"]) == c]
    if exact or len(c) < 4:
        return exact
    return [v for v in vehicles if _plate(v["plate_no"]).endswith(c)]


# ---------- 파일 읽기 ----------
class _Html(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows: list[list] = []
        self._row: list | None = None
        self._buf: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._buf = []
        elif tag == "br" and self._buf is not None:
            self._buf.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._buf is not None:
            self._row.append("".join(self._buf).strip() or None)
            self._buf = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._buf is not None:
            self._buf.append(data)


def _decode_html(data: bytes) -> str:
    m = re.search(rb"charset\s*=\s*[\"']?([A-Za-z0-9_-]+)", data[:4096])
    encs = ([m.group(1).decode("ascii")] if m else []) + ["utf-8-sig", "cp949"]
    for enc in encs:
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("cp949", errors="replace")


def file_kind(data: bytes) -> str:
    if data[:2] == b"PK":
        return "xlsx"
    if data[:8] == XLS_MAGIC:
        return "xls"
    if b"<table" in data[:65536].lower():
        return "html"
    raise ValueError("엑셀 파일이 아닙니다 (.xlsx · .xls만 올릴 수 있습니다)")


def _read_xlsx(data: bytes) -> list[list[list]]:
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        return [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
    finally:
        wb.close()


def _read_xls(data: bytes) -> list[list[list]]:
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError("구형 .xls를 읽는 모듈(xlrd)이 서버에 없습니다 — 관리자에게 알려 주세요") from exc
    book = xlrd.open_workbook(file_contents=data)
    out = []
    for sh in book.sheets():
        rows = []
        for i in range(sh.nrows):
            row = []
            for c in sh.row(i):
                if c.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(c.value, book.datemode))
                elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                else:
                    row.append(c.value)
            rows.append(row)
        out.append(rows)
    return out


def _read_html(data: bytes) -> list[list]:
    p = _Html()
    p.feed(_decode_html(data))
    p.close()
    return p.rows


def read_sheets(data: bytes) -> tuple[str, list[list[list]]]:
    kind = file_kind(data)
    try:
        if kind == "xlsx":
            sheets = _read_xlsx(data)
        elif kind == "xls":
            sheets = _read_xls(data)
        else:
            sheets = [_read_html(data)]
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001 — 손상 파일 등
        raise ValueError(f"파일을 읽을 수 없습니다 ({kind}) — 엑셀에서 열어 .xlsx로 다시 저장해 보세요") from exc
    return kind, [[[_cell(v) for v in r] for r in s] for s in sheets]


# ---------- 표 해석 ----------
def find_header(rows: list[list]) -> tuple[int, dict[str, int]]:
    for i, r in enumerate(rows[:HEADER_SCAN]):
        keys = [_key(v) for v in r]
        cols: dict[str, int] = {}
        for name, cands in COLS.items():
            for cand in cands:
                if cand in keys and keys.index(cand) not in cols.values():
                    cols[name] = keys.index(cand)
                    break
        if all(k in cols for k in REQUIRED):
            return i, cols
    top = " / ".join(t for r in rows[:3] for t in (_txt(v) for v in r) if t)[:150]
    raise NoHeader("머리글(일자·차량·합계)을 찾지 못했습니다" + (f" — 위쪽 내용: {top}" if top else ""))


def parse_rows(rows: list[list]) -> dict:
    h, cols = find_header(rows)

    def get(r, k):
        j = cols.get(k)
        return r[j] if j is not None and j < len(r) else None

    items: list[dict] = []
    skipped = 0
    subtotals: dict[str, int | None] = {}
    grand = None
    for i in range(h + 1, len(rows)):
        r = rows[i]
        if all(v is None for v in r):
            continue
        plate = _txt(get(r, "plate"))
        words = _key("".join(_txt(v) for v in r if isinstance(v, str)))
        if "합계" in plate:                      # 'NNNN 합계'
            subtotals[norm_code(plate.replace("합계", ""))] = _won(get(r, "total"))
            skipped += 1
            continue
        if "총합계" in words:                    # '[총 합계]'
            grand = _won(get(r, "total"))
            skipped += 1
            continue
        if not plate or "합계" in words:         # '(주)런치랩 합계' · '----' · 0만 있는 줄
            skipped += 1
            continue
        if len(items) >= MAX_ROWS:
            raise ValueError(f"주유 행이 {MAX_ROWS}행을 넘습니다 — 파일을 나눠 올리세요")
        product = _txt(get(r, "product")) or None
        items.append({
            "row": i + 1, "code": norm_code(plate),
            "date": parse_date(get(r, "date")), "date_raw": _txt(get(r, "date")),
            "product": product, "fuel_type": fuel_kind(product),
            "quantity": _num(get(r, "qty")),
            "supply": _won(get(r, "supply")), "vat": _won(get(r, "vat")), "total": _won(get(r, "total")),
        })

    checks: list[dict] = []
    by_code: dict[str, int] = {}
    for it in items:
        by_code[it["code"]] = by_code.get(it["code"], 0) + (it["total"] or 0)
    for code, sub in subtotals.items():
        got = by_code.get(code, 0)
        if sub is not None and sub != got:
            checks.append({"level": "warn", "text": f"차량 {code} 합계 줄 {sub:,}원 ≠ 행 합계 {got:,}원 — 빠진 행이 없는지 확인하세요"})
    total = sum(by_code.values())
    if grand is not None:
        if grand == total:
            checks.append({"level": "ok", "text": f"파일 총 합계 {grand:,}원과 행 합계가 같습니다 ({len(items)}행)"})
        else:
            checks.append({"level": "warn", "text": f"파일 총 합계 {grand:,}원 ≠ 행 합계 {total:,}원 — 빠진 행이 없는지 확인하세요"})
    return {"items": items, "skipped": skipped, "checks": checks, "header_row": h + 1}


def parse_file(data: bytes) -> dict:
    kind, sheets = read_sheets(data)
    err: Exception | None = None
    for rows in sheets:
        try:
            out = parse_rows(rows)
        except NoHeader as exc:
            err = exc
            continue
        out["file_kind"] = kind
        return out
    raise ValueError(str(err) if err else "빈 파일입니다")


# ---------- 행 판정 ----------
def pair_refunds(items: list[dict]) -> None:
    """취소(음수) 행 ↔ 같은 차량·같은 금액 원 거래 (같은 날 우선, 없으면 가장 가까운 날). 짝이면 둘 다 'pair'=상대 행."""
    for it in items:
        if it.get("pair") or it["total"] is None or it["total"] >= 0:
            continue
        best, gap_best = None, None
        for o in items:
            if o is it or o.get("pair") or o["code"] != it["code"] or o["total"] != -it["total"]:
                continue
            gap = abs((o["date"] - it["date"]).days) if o["date"] and it["date"] else 9999
            if best is None or gap < gap_best:
                best, gap_best = o, gap
        if best is not None:
            it["pair"], best["pair"] = best["row"], it["row"]


def judge(items: list[dict], vehicles: list[dict], vehicle_map: dict[str, int],
          existing: set, closed: set, today_: date) -> list[dict]:
    """vehicles: [{id, plate_no, fuel_type, vat_deductible}], vehicle_map: {끝번호: 차량 id} (직접 선택),
    existing: {(차량 id, 날짜, 합계)} DB 기록, closed: {마감된 달의 1일}.
    level: ok · warn(선택 가능) · need_vehicle · error · cancel(취소 짝)."""
    pair_refunds(items)
    by_id = {int(v["id"]): v for v in vehicles}
    seen: set = set()
    out: list[dict] = []
    for it in items:
        r = {**it, "vehicle_id": None, "plate_no": None, "vehicle_fuel": None, "candidates": [], "auto": False,
             "level": "ok", "selectable": False, "checked": False, "messages": [],
             "total_amount": it["total"], "supply_amount": it["supply"], "vat_amount": it["vat"],
             "unit_price": None, "cost": None}
        out.append(r)
        msgs = r["messages"]
        if it.get("pair"):
            r["level"] = "cancel"
            msgs.append(f"취소 거래와 짝({it['pair']}행) — 둘 다 저장하지 않습니다")
            continue

        cands = candidates(it["code"], vehicles)
        r["candidates"] = [int(v["id"]) for v in cands]
        v = None
        chosen = vehicle_map.get(it["code"])
        if chosen is not None:
            v = by_id.get(int(chosen))
            if v is None:
                msgs.append("선택한 차량을 찾을 수 없습니다 — 다시 선택하세요")
        elif len(cands) == 1:
            v, r["auto"] = cands[0], True
        if v is not None:
            r.update(vehicle_id=int(v["id"]), plate_no=v["plate_no"], vehicle_fuel=v["fuel_type"])

        d, total = it["date"], it["total"]
        errors = []
        if d is None:
            errors.append(f"날짜를 읽을 수 없습니다 ({it['date_raw'] or '빈 칸'})")
        elif d > today_:
            errors.append("미래 날짜입니다")
        elif d.replace(day=1) in closed:
            errors.append(f"{d.year}년 {d.month}월 유류비는 마감됐습니다 — 마감 해제 후 올리세요")
        if not total:
            errors.append("합계 금액이 없습니다")
        elif total < 0:
            errors.append("취소(음수) 거래인데 짝이 되는 원 거래가 파일에 없습니다 — 이미 저장된 기록이면 유류비 탭에서 직접 삭제하세요")
        if errors:
            r["level"] = "error"
            msgs.extend(errors)
            continue

        if v is None:
            r["level"] = "need_vehicle"
            if chosen is None:
                msgs.append(f"차량 끝번호 {it['code']}와 맞는 차량이 없습니다 — 차량을 선택하세요" if not cands else
                            f"차량 끝번호 {it['code']}가 {len(cands)}대와 맞습니다 "
                            f"({', '.join(c['plate_no'] for c in cands)}) — 차량을 선택하세요")
            continue

        warn_off, warn_on = [], []
        try:
            tot, sup, vat = amounts(total, it["supply"], it["vat"])
        except ValueError:
            tot, sup, vat = amounts(total)
            warn_on.append(f"공급가액+부가세가 합계 {total:,}원과 달라 합계 ÷ 1.1로 다시 계산했습니다")
        r.update(total_amount=tot, supply_amount=sup, vat_amount=vat,
                 cost=sup if v.get("vat_deductible") else tot)
        q = it["quantity"]
        if q is not None and 0 < q <= 100000:
            r["quantity"] = round(q, 3)
            up = round(tot / q, 2)
            r["unit_price"] = up if 0 < up <= 10000 else None
        else:
            r["quantity"] = None

        ft = it["fuel_type"]
        if ft and ft != v["fuel_type"]:
            warn_off.append(f"제품명 '{it['product']}'({FUEL_LABEL[ft]})이 차량 연료"
                            f"({FUEL_LABEL.get(v['fuel_type'], v['fuel_type'])})와 다릅니다 — 저장하면 차량 연료로 기록됩니다")
        key = (r["vehicle_id"], d, tot)
        if key in existing:
            warn_off.append("같은 차량·날짜·금액 기록이 이미 있습니다 (중복 의심)")
        elif key in seen:
            warn_off.append("파일 안에 같은 차량·날짜·금액 행이 또 있습니다 (중복 의심)")
        seen.add(key)

        msgs.extend(warn_off + warn_on)
        r["level"] = "warn" if msgs else "ok"
        r["selectable"] = True
        r["checked"] = not warn_off
    return out
