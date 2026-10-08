"""S2-P4a-2 근무표 엑셀 생성·검증 (DB 없음 → 단위 테스트 대상).

- 시스템 근무 기록이 원본, 이 파일은 제출용 결과물. 양식은 코드로 생성 (공개 레포에 바이너리·실명 없음)
- 초과수당 미포함 (2026-10-07). 월급여 = 시급 × 월 유급시간, 총지급액 = 월급여 + 인센티브 + 주말수당
- 검증: 만든 파일을 parse_timesheet로 다시 읽어(수식 칸 제외) 출근·퇴근·휴게·시급·인원을 기록과 1분 단위 대조
"""
from __future__ import annotations

import calendar
import io
from datetime import date, datetime, time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from app.services.timesheet_excel import LABELS, hm, parse_timesheet

WEEKDAYS = "월화수목금토일"
SHEET_TITLE = {"labor": "근로소득", "business": "사업소득"}
HDR = 5          # 날짜 행
FIRST_ROW = 8    # 첫 사람 출근 행
DAY_COL = 5      # E열 = 1일
MONEY = "#,##0"

_THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
F_SAT = PatternFill("solid", fgColor="DDEBF7")
F_SUN = PatternFill("solid", fgColor="FCE4D6")
F_HEAD = PatternFill("solid", fgColor="F2F2F2")
F_SUM = PatternFill("solid", fgColor="FFF2CC")


def _t(m: int | None) -> time | None:
    return None if m is None else time(m // 60, m % 60)


def _cols(n: int, sheet: str) -> dict:
    cm = DAY_COL + n
    cp = cm + 2
    c = {"month": cm, "pay": cp, "inc": cp + 1, "wk": cp + 2, "total": cp + 3,
         "vat": cp + 4 if sheet == "business" else None}
    c["last"] = c["vat"] or c["total"]
    return c


def schedule_note(spans: list[tuple[date, str | None, str | None, int]]) -> str:
    """'시간대' 칸. spans = [(시작일, 지정 출근, 지정 퇴근, 시급)] 시작일 순.
    '07:00~12:30 / 9/8~ 07:00~13:30', 시급만 바뀌면 '9/15~ 시급 17,000'"""
    parts: list[str] = []
    prev: tuple[str, int] | None = None
    for frm, ws, we, amt in spans:
        t = f"{ws}~{we}"
        head = f"{frm.month}/{frm.day}~ "
        if prev is None:
            parts.append(t)
        else:
            if t != prev[0]:
                parts.append(head + t)
            if amt != prev[1]:
                parts.append(head + f"시급 {amt:,}")
        prev = (t, amt)
    return " / ".join(parts)


def _person(ws, r: int, p: dict, n: int, c: dict, fills: dict) -> None:
    L = get_column_letter
    first_l, last_l = L(DAY_COL), L(DAY_COL + n - 1)
    sr = r + 3
    ws.cell(r, 1, p["name"])
    ws.cell(r, 2, int(p["rate"])).number_format = MONEY
    if p.get("note"):
        ws.cell(r, 3, p["note"])
    for k, lab in enumerate(LABELS):
        ws.cell(r + k, 4, lab)
    for d in range(1, n + 1):
        col = DAY_COL + d - 1
        cl = L(col)
        v = p["days"].get(d)
        if v:
            ws.cell(r, col, _t(v["in"]))
            ws.cell(r + 1, col, _t(v["out"]))
            if v.get("break"):
                ws.cell(r + 2, col, _t(v["break"]))
        for k in range(3):
            ws.cell(r + k, col).number_format = "h:mm"
        ws.cell(sr, col, f'=IF(COUNT({cl}{r}:{cl}{r + 1})=2,{cl}{r + 1}-{cl}{r}-N({cl}{r + 2}),"")').number_format = "[h]:mm"
    m_l, p_l, w_l, t_l = (L(c[k]) for k in ("month", "pay", "wk", "total"))
    ws.cell(sr, c["month"], f"=SUM({first_l}{sr}:{last_l}{sr})").number_format = "[h]:mm"
    ws.cell(r, c["pay"], f"={m_l}{sr}*24").number_format = "0.00"
    pv = p.get("pay_value")
    ws.cell(sr, c["pay"], int(pv) if pv is not None else f"=ROUND($B{r}*{p_l}{r},0)").number_format = MONEY
    for key, col in (("incentive", c["inc"]), ("weekend", c["wk"])):
        if p.get(key):
            ws.cell(sr, col, int(p[key]))
        ws.cell(sr, col).number_format = MONEY
    ws.cell(sr, c["total"], f"=SUM({p_l}{sr}:{w_l}{sr})").number_format = MONEY
    if c["vat"] and p.get("vat"):
        ws.cell(sr, c["vat"], f"=ROUND({t_l}{sr}*1.1,0)").number_format = MONEY
    for k in range(4):
        for col in range(1, c["last"] + 1):
            x = ws.cell(r + k, col)
            x.border, x.alignment = BOX, CENTER
            f = fills.get(col)
            if f:
                x.fill = f
            elif k == 3 and col >= 4:
                x.fill = F_SUM
    for col in (1, 2, 3):                       # 스타일 적용 후 병합 (병합 칸에 값 쓰기 방지)
        ws.merge_cells(start_row=r, start_column=col, end_row=sr, end_column=col)


def build_timesheet(sheet: str, month: date, people: list[dict], *, factory: str = "", dept: str = "") -> bytes:
    """people = [{name, rate, note, days: {일: {in, out, break}(분)}, pay_value, incentive, weekend, vat}]"""
    if sheet not in SHEET_TITLE:
        raise ValueError("근무표 구분은 labor / business 입니다")
    first = month.replace(day=1)
    n = calendar.monthrange(first.year, first.month)[1]
    c = _cols(n, sheet)
    L = get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = f"{first.year}.{first.month:02d}"
    for r, (k, v) in enumerate((("공장", factory), ("부서", dept), ("연월", datetime(first.year, first.month, 1))), 1):
        ws.cell(r, 1, k).font = BOLD
        ws.cell(r, 2, v)
    ws.cell(3, 2).number_format = "yyyy-mm"
    ws.cell(1, 4, f"{SHEET_TITLE[sheet]} 근무시간표").font = Font(bold=True, size=14)

    fills: dict[int, PatternFill | None] = {}
    for d in range(1, n + 1):
        col = DAY_COL + d - 1
        wd = date(first.year, first.month, d).weekday()
        fills[col] = F_SAT if wd == 5 else F_SUN if wd == 6 else None
        ws.cell(HDR, col, d)
        ws.cell(HDR + 1, col, WEEKDAYS[wd])
    heads = {c["month"]: "월", c["pay"]: "월급여", c["inc"]: "인센티브", c["wk"]: "주말특별수당", c["total"]: "총지급액"}
    if c["vat"]:
        heads[c["vat"]] = "부가세 포함"
    for col, v in heads.items():
        ws.cell(HDR, col, v)
    ws.cell(HDR + 1, c["month"], "합계")
    for col, v in enumerate(("성명", "시급", "시간대", "구분"), 1):
        ws.cell(HDR + 2, col, v)
    for r in range(HDR, HDR + 3):
        for col in range(1, c["last"] + 1):
            x = ws.cell(r, col)
            x.font, x.alignment, x.border = BOLD, CENTER, BOX
            x.fill = fills.get(col) or F_HEAD

    r = FIRST_ROW
    for p in people:
        _person(ws, r, p, n, c, fills)
        r += 4
    if people:
        ws.cell(r, 1, "총계").font = BOLD
        for key in ("pay", "inc", "wk", "total", "vat"):
            col = c[key]
            if col is None:
                continue
            cl = L(col)
            x = ws.cell(r, col, f'=SUMIF($D${FIRST_ROW}:$D${r - 1},"합계",{cl}{FIRST_ROW}:{cl}{r - 1})')
            x.number_format, x.font, x.border = MONEY, BOLD, BOX

    for col, w in ((1, 10), (2, 9), (3, 20), (4, 6)):
        ws.column_dimensions[L(col)].width = w
    for col in range(DAY_COL, DAY_COL + n):
        ws.column_dimensions[L(col)].width = 6.5
    ws.column_dimensions[L(c["month"])].width = 8
    ws.column_dimensions[L(c["month"] + 1)].width = 2
    for col in range(c["pay"], c["last"] + 1):
        ws.column_dimensions[L(col)].width = 12
    ws.freeze_panes = ws.cell(FIRST_ROW, DAY_COL)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _where(x: dict) -> str:
    s = " ".join(v for v in (x.get("name") or "", f"{x['day']}일" if x.get("day") else "") if v)
    return f"{s}: {x['msg']}" if s else x["msg"]


def verify_timesheet(data: bytes, people: list[dict]) -> list[str]:
    """만든 파일 재검증 → 오류 목록 (비어 있어야 다운로드). 수식 칸은 비우고 읽음(엑셀 계산 전이라 값 없음)."""
    ws = load_workbook(io.BytesIO(data)).active
    rows = [[None if isinstance(v, str) and v.startswith("=") else v for v in row]
            for row in ws.iter_rows(values_only=True)]
    res = parse_timesheet(rows, ws.title)
    errs = [_where(x) for x in res["issues"] if x["level"] == "error"]
    names = [p["name"] for p in res["people"]]
    if names != [p["name"] for p in people]:
        errs.append(f"인원·순서가 기록과 다릅니다 (파일 {len(names)}명, 기록 {len(people)}명)")
    got = {p["name"]: p for p in res["people"]}
    for p in people:
        g = got.get(p["name"])
        if g is None:
            continue
        if int(g["rate"] or 0) != int(p["rate"]):
            errs.append(f"{p['name']}: 시급 칸이 기록과 다릅니다")
        want = {d: (hm(v["in"]), hm(v["out"]), v.get("break") or None) for d, v in p["days"].items()}
        have = {d: (v["in"], v["out"], v["break_min"] or None) for d, v in g["days"].items()}
        for d in sorted(set(want) | set(have)):
            if want.get(d) != have.get(d):
                errs.append(f"{p['name']} {d}일: 파일 {have.get(d)} ≠ 기록 {want.get(d)}")
    return errs
