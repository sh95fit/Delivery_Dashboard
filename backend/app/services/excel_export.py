"""S3-P1b-2b-3 엑셀 내보내기 공통 (openpyxl).
모든 비용·매출 화면이 같은 모양(굵은 머리글·고정·필터·원 단위 서식·합계 줄)으로 내려받도록 한 곳에서 만든다.
sheet = {"title": 시트명, "columns": [(머리글, 키, 형식)], "rows": [dict], "totals": dict | None}
형식: won(원) · qty(수량 소수 3자리) · num · date · text
"""
from __future__ import annotations

import re
from datetime import date
from io import BytesIO
from urllib.parse import quote

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FMT = {"won": "#,##0", "qty": "#,##0.000", "num": "#,##0.##", "date": "yyyy-mm-dd"}
WIDTH = {"won": 13, "qty": 12, "num": 10, "date": 12, "text": 16}


def _title(t) -> str:
    return re.sub(r"[\[\]:*?/\\]", " ", str(t))[:31] or "Sheet"


def _val(v, fmt: str):
    if v is None:
        return None
    if fmt == "date" and isinstance(v, str):
        try:
            return date.fromisoformat(v[:10])
        except ValueError:
            return v
    if isinstance(v, bool):
        return "예" if v else "아니오"
    if isinstance(v, (list, tuple, set)):
        return ", ".join(str(x) for x in v)
    return v


def build_xlsx(sheets: list[dict]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="EEF2F7")
    for s in sheets:
        ws = wb.create_sheet(_title(s["title"]))
        cols = s["columns"]
        ws.append([c[0] for c in cols])
        for cell in ws[1]:
            cell.font, cell.fill = bold, fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for r in s["rows"]:
            ws.append([_val(r.get(c[1]), c[2]) for c in cols])
        last = get_column_letter(len(cols))
        ws.auto_filter.ref = f"A1:{last}{ws.max_row}"
        if s.get("totals"):
            ws.append([_val(s["totals"].get(c[1]), c[2]) for c in cols])
            for cell in ws[ws.max_row]:
                cell.font = bold
        for j, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(j)].width = max(WIDTH.get(c[2], 14), len(c[0]) * 2 + 2)
            fmt = FMT.get(c[2])
            if fmt:
                for (cell,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    cell.number_format = fmt
        ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(filename: str, sheets: list[dict]) -> Response:
    body = build_xlsx(sheets)
    ascii_name = re.sub(r"[^A-Za-z0-9._-]", "_", filename) or "export.xlsx"
    disp = f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"
    return Response(body, media_type=XLSX, headers={"Content-Disposition": disp, "Cache-Control": "no-store"})
